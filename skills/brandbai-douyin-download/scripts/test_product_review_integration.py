"""Offline integration contracts; never connects to a real platform."""
import copy
import json
import shutil
import unittest
import uuid
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook
import download_creator_works as works
import run_foundation as foundation
import build_foundation_workbooks as books
from browser_collect_product_reviews import ProductReviewError, normalize_review

WORK = "7000000000000000001"
URL = "https://www.douyin.com/video/" + WORK
PRODUCT = {"title": "合成商品", "shop_name": "合成店铺", "product_id": "90000001"}


@contextmanager
def temporary():
    path = Path(__file__).parent / ("_product_integration_" + uuid.uuid4().hex)
    path.mkdir()
    try:
        yield path
    finally:
        shutil.rmtree(path)


def arguments(path, *extra, capability="all"):
    return foundation.build_parser().parse_args([capability, "--video", URL, "--assets", "none",
        "--profile-dir", str(path / "profile"), "--out", str(path / "delivery"),
        "--product-reviews", *extra])


def review_manifest(reason="run_budget", complete=False):
    return {"task_id": "synthetic-task", "source_work_id": WORK, "privacy_mode": "hash", "product": copy.deepcopy(PRODUCT),
        "review_count": 1, "status": "complete" if complete else "partial", "done_reason": reason,
        "completeness": "complete_visible_panel_exhausted" if complete else "partial_" + reason}


def save_review(path, reason="run_budget", complete=False):
    path.mkdir(parents=True, exist_ok=True)
    manifest = review_manifest(reason, complete)
    row = normalize_review({"reviewId": "12345678", "reviewerName": "合成评价者", "dateText": "4个月前",
        "content": "商品评价独立正文", "purchasedSku": "规格A", "images": []}, manifest)
    (path / "product_review_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    (path / "product_reviews.jsonl").write_text(json.dumps(row, ensure_ascii=False) + "\n", encoding="utf-8")
    return manifest, row


class ProductReviewIntegrationTests(unittest.TestCase):
    def test_opt_in_is_not_a_default_and_dry_plan_is_bounded(self):
        with temporary() as path:
            args = arguments(path)
            plan = foundation.all_plan(args)
            self.assertTrue(args.commerce_detail)
            self.assertEqual(plan["product_reviews"]["max_reviews"], 200)
            self.assertIn("05_商品评价.xlsx", plan["ordinary_files"])
            self.assertIn("data/商品评价", plan["raw_data"])
            ordinary = foundation.build_parser().parse_args(["all", "--video", URL, "--profile-dir", "profile", "--out", "out"])
            self.assertFalse(ordinary.product_reviews)
            self.assertNotIn("05_商品评价.xlsx", foundation.all_plan(ordinary)["ordinary_files"])

    def test_creator_multiwork_and_invalid_budgets_fail_before_browser(self):
        with temporary() as path:
            invalid = [arguments(path, "--video", "https://www.douyin.com/video/7000000000000000002"),
                arguments(path, "--max-product-reviews", "0"), arguments(path, "--product-review-max-scrolls", "-1"),
                arguments(path, "--product-review-max-seconds", "nan")]
            invalid.append(foundation.build_parser().parse_args(["all", "--creator", "https://www.douyin.com/user/test", "--product-reviews", "--profile-dir", "p", "--out", "o"]))
            for args in invalid:
                with self.subTest(args=args), self.assertRaises(foundation.FoundationError):
                    foundation.all_plan(args)

    def test_child_command_forwards_explicit_budget_privacy_and_resume(self):
        with temporary() as path:
            args = arguments(path, "--resume", "--privacy-mode", "raw", "--max-product-reviews", "35", capability="works")
            command = foundation.child_command(args, Path("scripts"))
            for value in ("--commerce-detail", "--product-reviews", "--resume", "raw", "35"):
                self.assertIn(value, command)

    def test_partial_stage_uses_same_page_and_restores_only_same_work(self):
        with temporary() as path:
            args = arguments(path)
            page = Mock()
            def collector(actual_page, work_id, product, out, **options):
                self.assertIs(actual_page, page)
                self.assertEqual(work_id, WORK)
                self.assertEqual(product, PRODUCT)
                self.assertEqual(options, dict(max_reviews=200, max_scrolls=200, max_seconds=600.0, privacy_mode="hash", resume=False))
                save_review(out)
                return 3
            with patch("browser_collect_product_reviews.collect_product_reviews", side_effect=collector):
                self.assertEqual(works.run_product_reviews_on_page(page, args, {"aweme_id": WORK}, PRODUCT), 3)
            self.assertTrue(args._product_review_summary["work_comments_safe"])
            page.goto.assert_called_once_with(URL, wait_until="domcontentloaded", timeout=60_000)

    def test_source_warning_does_not_navigate_into_work_comment_stage(self):
        with temporary() as path:
            args = arguments(path); page = Mock()
            def collector(_page, _work_id, _product, out, **_options):
                save_review(out, "surface_or_filter_changed")
                return 3
            with patch("browser_collect_product_reviews.collect_product_reviews", side_effect=collector):
                self.assertEqual(works.run_product_reviews_on_page(page, args, {"aweme_id": WORK}, PRODUCT), 3)
            page.goto.assert_not_called()
            self.assertFalse(args._product_review_summary["work_comments_safe"])

    def test_code_zero_without_confirmed_manifest_remains_partial(self):
        with temporary() as path:
            args = arguments(path); page = Mock()
            with patch("browser_collect_product_reviews.collect_product_reviews", return_value=0):
                self.assertEqual(works.run_product_reviews_on_page(page, args, {"aweme_id": WORK}, PRODUCT), 3)
            page.goto.assert_not_called()

    def test_user_pause_never_navigates_or_continues_work_comments(self):
        with temporary() as path:
            args = arguments(path); page = Mock()
            def collector(_page, _work, _product, out, **_options):
                save_review(out, "user_paused")
                return 3
            with patch("browser_collect_product_reviews.collect_product_reviews", side_effect=collector):
                self.assertEqual(works.run_product_reviews_on_page(page, args, {"aweme_id": WORK}, PRODUCT), 3)
            page.goto.assert_not_called()
            self.assertFalse(args._product_review_summary["work_comments_safe"])

    def test_resume_rejection_preserves_old_rows_and_partial_attempt(self):
        with temporary() as path:
            args = arguments(path, "--resume"); page = Mock()
            raw = works.product_review_output_dir(args)
            save_review(raw, "source_exhausted", True)
            original = (raw / "product_reviews.jsonl").read_bytes()
            with patch("browser_collect_product_reviews.collect_product_reviews", side_effect=ProductReviewError("unsafe browser detail not exported")):
                self.assertEqual(works.run_product_reviews_on_page(page, args, {"aweme_id": WORK}, PRODUCT), 3)
            self.assertEqual((raw / "product_reviews.jsonl").read_bytes(), original)
            self.assertEqual(args._product_review_summary["status"], "partial")
            self.assertNotIn("unsafe browser", json.dumps(args._product_review_summary))
            page.goto.assert_not_called()

    def test_resume_identity_includes_product_request_and_privacy(self):
        with temporary() as path:
            args = arguments(path); works.validate_product_review_args(args)
            identity = foundation.work_input_identity(args)
            self.assertEqual(identity, works.input_identity(args))
            manifest = path / "download_manifest.json"; payload = path / "works.json"
            manifest.write_text(json.dumps({"status": "complete", "works_selected": 1, "input_identity": identity}), encoding="utf-8")
            payload.write_text(json.dumps([{"source_url": URL}]), encoding="utf-8")
            foundation.validate_resume_works(manifest, payload, "", 5, identity)
            args.privacy_mode = "raw"
            with self.assertRaises(foundation.FoundationError):
                foundation.validate_resume_works(manifest, payload, "", 5, foundation.work_input_identity(args))
            args.privacy_mode = "hash"; args.product_reviews = False
            with self.assertRaises(foundation.FoundationError):
                foundation.validate_resume_works(manifest, payload, "", 5, foundation.work_input_identity(args))

    def test_resume_retries_product_stage_without_touching_saved_assets(self):
        with temporary() as path:
            args = arguments(path, "--resume"); works.validate_product_review_args(args)
            output = Path(args.out); output.mkdir()
            saved_work = [{"aweme_id": WORK, "type": "视频", "source_url": URL}]
            (output / "works.json").write_text(json.dumps(saved_work), encoding="utf-8")
            (output / "download_manifest.json").write_text(json.dumps({"status": "complete", "input_identity": works.input_identity(args), "product_reviews": {"status": "partial"}}), encoding="utf-8")
            original = (output / "works.json").read_bytes(); page = Mock()
            def stage(actual_page, actual_args, actual_work, product):
                self.assertIs(actual_page, page); self.assertEqual(actual_work, saved_work[0]); self.assertEqual(product, PRODUCT)
                actual_args._product_review_summary = {"requested": True, "status": "partial", "exit_code": 3, "work_comments_safe": False}
                return 3
            with patch.object(works, "collect_visible_commerce_detail", return_value={"products": [PRODUCT]}), patch.object(works, "run_product_reviews_on_page", side_effect=stage) as collector:
                self.assertEqual(works.resume_product_review_stage(SimpleNamespace(pages=[page]), args), 3)
            collector.assert_called_once()
            self.assertEqual((output / "works.json").read_bytes(), original)

    def test_shared_context_blocks_comments_after_unsafe_product_result(self):
        with temporary() as path:
            args = arguments(path); context = Mock()
            @contextmanager
            def browser():
                yield SimpleNamespace(chromium=SimpleNamespace(launch_persistent_context=lambda **_kw: context))
            def run_works(actual_args, **_kw):
                actual_args._product_review_summary = {"exit_code": 3, "work_comments_safe": False}
                return 3
            comments = Mock(return_value=0)
            result = foundation.run_shared_browser_stages(args, SimpleNamespace(), path / "trace.jsonl",
                playwright_factory=browser, works_runner=run_works, comments_runner=comments, chrome_finder=lambda _: "chrome")
            self.assertEqual(result, (3, 0)); comments.assert_not_called(); context.close.assert_called_once()
            self.assertEqual(args._comments_skipped_reason, "product_review_requires_attention")

    def test_complete_unknown_id_evaluations_are_reused_without_reopening_product(self):
        with temporary() as path:
            args = arguments(path, "--resume"); works.validate_product_review_args(args)
            output = Path(args.out); output.mkdir()
            product = {**PRODUCT, "product_id": ""}
            work = {"aweme_id": WORK, "type": "视频", "source_url": URL, "commerce": {"products": [product]}}
            (output / "works.json").write_text(json.dumps([work]), encoding="utf-8")
            (output / "download_manifest.json").write_text(json.dumps({"status": "complete", "input_identity": works.input_identity(args)}), encoding="utf-8")
            raw = works.product_review_output_dir(args); saved, row = save_review(raw, "source_exhausted", True)
            saved["product"] = product; row["product_id"] = ""
            (raw / "product_review_manifest.json").write_text(json.dumps(saved), encoding="utf-8")
            (raw / "product_reviews.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
            page = Mock()
            with patch.object(works, "collect_visible_commerce_detail") as detail, patch.object(works, "run_product_reviews_on_page") as collect:
                self.assertEqual(works.resume_product_review_stage(SimpleNamespace(pages=[page]), args), 0)
            detail.assert_not_called(); collect.assert_not_called()
            self.assertTrue(args._product_review_summary["reused_complete"])
            self.assertTrue(args._product_review_summary["work_comments_safe"])
            saved["review_count"] = 2
            (raw / "product_review_manifest.json").write_text(json.dumps(saved), encoding="utf-8")
            with self.assertRaisesRegex(works.WorkDownloadError, "count"):
                works.completed_product_review_summary(args, work)

    def test_ordinary_product_book_is_separate_and_preserves_partial(self):
        with temporary() as path:
            raw = path / "商品评价"; save_review(raw)
            output = path / "05_商品评价.xlsx"
            manifest = {"product_reviews": {"requested": True, "privacy_mode": "hash", "status": "partial", "exit_code": 3, "done_reason": "run_budget"}}
            summary = books.build_product_review_delivery(raw, output, [{"aweme_id": WORK}], manifest)
            self.assertEqual(summary["exit_code"], 3); self.assertEqual(summary["review_count"], 1)
            book = load_workbook(output)
            self.assertEqual(book.sheetnames, ["采集说明", "商品评价"])
            self.assertEqual(book["商品评价"]["A2"].value, "商品评价独立正文")
            self.assertEqual(book["商品评价"]["N2"].value, WORK); book.close()
            books.build_explanation(path, "合成", 1, 0, 0, manifest, {"status": "not_requested"}, summary)
            self.assertIn("05_商品评价.xlsx", (path / "04_采集说明.md").read_text(encoding="utf-8"))
            self.assertIn("部分完成", (path / "04_采集说明.md").read_text(encoding="utf-8"))

    def test_missing_product_output_is_not_a_confirmed_empty_list(self):
        with temporary() as path:
            summary = books.build_product_review_delivery(path / "missing", path / "05.xlsx", [{"aweme_id": WORK}],
                {"product_reviews": {"requested": True, "privacy_mode": "hash", "exit_code": 3}})
            self.assertEqual(summary["exit_code"], 3)
            self.assertNotEqual(summary["completeness"], "complete_visible_panel_exhausted")

    def test_foreign_product_rows_are_rejected_without_overwriting_export(self):
        with temporary() as path:
            raw = path / "raw"; save_review(raw)
            output = path / "05.xlsx"; output.write_bytes(b"previous-export")
            manifest = json.loads((raw / "product_review_manifest.json").read_text(encoding="utf-8"))
            manifest["privacy_mode"] = "raw"
            (raw / "product_review_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            with self.assertRaises(ValueError):
                books.build_product_review_delivery(raw, output, [{"aweme_id": WORK}], {"product_reviews": {"privacy_mode": "hash"}})
            self.assertEqual(output.read_bytes(), b"previous-export")

    def test_full_offline_delivery_keeps_three_datasets_and_partial_exit(self):
        with temporary() as path:
            args = arguments(path, "--zip")
            events = []
            page = Mock()
            context = SimpleNamespace(pages=[page], close=Mock())
            @contextmanager
            def browser():
                yield SimpleNamespace(chromium=SimpleNamespace(launch_persistent_context=lambda **_kw: context))
            def product_detail(actual_page, work_id):
                self.assertIs(actual_page, page); self.assertEqual(work_id, WORK)
                events.append("product_detail")
                return {"status": "detail_observed", "source_work_id": WORK, "products": [copy.deepcopy(PRODUCT)]}
            def product_reviews(actual_page, work_id, _product, out, **_options):
                self.assertIs(actual_page, page); self.assertEqual(work_id, WORK)
                events.append("product_reviews"); save_review(out)
                return 3
            def comments(actual_args, browser_context):
                self.assertIs(browser_context, context)
                self.assertEqual(page.goto.call_args.args[0], URL)
                events.append("work_comments")
                output = Path(actual_args.out); output.mkdir(parents=True)
                (output / "comments.csv").write_text("aweme_id,comment_id,text,reply_level\n" + WORK + ",70000001,作品评论独立正文,0\n", encoding="utf-8")
                (output / "run_manifest.json").write_text(json.dumps({"status": "complete_source_visible", "privacy_mode": "hash"}), encoding="utf-8")
                return 0
            def stages(work_args, comment_args, trace_path, **options):
                return foundation.run_shared_browser_stages(work_args, comment_args, trace_path,
                    playwright_factory=browser, works_runner=works.run, comments_runner=comments,
                    chrome_finder=lambda _: "chrome", **options)
            def build(command, **_options):
                return SimpleNamespace(returncode=books.main(command[2:]))
            with patch.object(works, "collect_visible_commerce_anchor", return_value={"status": "visible_name_only"}), \
                 patch.object(works, "collect_visible_commerce_detail", side_effect=product_detail), \
                 patch("browser_collect_product_reviews.collect_product_reviews", side_effect=product_reviews):
                self.assertEqual(foundation.run_all(args, runner=build, browser_stage_runner=stages), 3)
            self.assertEqual(events, ["product_detail", "product_reviews", "work_comments"])
            output = Path(args.out)
            self.assertTrue((output / "data" / "作品采集" / "works.json").is_file())
            self.assertTrue((output / "data" / "商品评价" / "product_reviews.jsonl").is_file())
            delivery = json.loads((output / "data" / "foundation_manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(delivery["status"], "partial")
            work_manifest = json.loads((output / "data" / "作品采集" / "download_manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(work_manifest["status"], "complete")
            self.assertEqual(work_manifest["product_reviews"]["exit_code"], 3)
            for filename, expected, excluded in [("02_评论明细.xlsx", "作品评论独立正文", "商品评价独立正文"),
                    ("05_商品评价.xlsx", "商品评价独立正文", "作品评论独立正文")]:
                book = load_workbook(output / filename)
                content = "\n".join(str(cell.value) for sheet in book for row in sheet for cell in row)
                self.assertIn(expected, content); self.assertNotIn(excluded, content); book.close()
            self.assertTrue(output.with_suffix(".zip").is_file())

    def test_failed_review_workbook_never_repackages_existing_old_files(self):
        with temporary() as path:
            args = arguments(path, "--zip", "--skip-comments")
            output = Path(args.out); output.mkdir()
            for name in ("01_作品清单.xlsx", "02_评论明细.xlsx", "04_采集说明.md", "05_商品评价.xlsx"):
                (output / name).write_bytes(b"previous-attempt")
            archive = output.with_suffix(".zip"); archive.write_bytes(b"previous-zip")
            def stages(work_args, _comment_args, _trace_path, **_options):
                raw = Path(work_args.out); raw.mkdir(parents=True)
                (raw / "works.json").write_text(json.dumps([{"aweme_id": WORK}]), encoding="utf-8")
                work_args._product_review_summary = {"requested": True, "status": "partial", "exit_code": 3}
                return 3, 0
            with patch.object(foundation, "package_directory") as package:
                self.assertEqual(foundation.run_all(args, runner=lambda *_a, **_k: SimpleNamespace(returncode=1), browser_stage_runner=stages), 3)
                package.assert_not_called()
            saved = json.loads((output / "data" / "foundation_manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(saved["status"], "partial")
            self.assertEqual(saved["error"], "product_review_workbook_build_failed")
            self.assertEqual(saved["package_status"], "not_created_this_attempt")
            self.assertEqual(archive.read_bytes(), b"previous-zip")
            self.assertEqual((output / "05_商品评价.xlsx").read_bytes(), b"previous-attempt")


if __name__ == "__main__":
    unittest.main()
