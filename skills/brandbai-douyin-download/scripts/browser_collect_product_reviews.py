"""Bounded product reviews, independent from work comments (no platform API calls).

The page collector is vendored from BrandBAI extension v0.11.156. Python owns
durable acknowledgements, anonymization, same-source resume and local exports.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

ROOT = Path(__file__).resolve().parent
BRIDGE = "globalThis.__brandbaiSkillProductReviewsV060"
LEASE_FIELDS = ("documentToken", "contextKey", "generation", "sourceWorkId", "sourceSurfaceInstance",
                "productPanelInstanceId", "reviewSurfaceInstanceId", "filterKey")
COMPLETE = "complete_visible_panel_exhausted"
ALLOWED_HOSTS = ("ecombdimg.com", "byteimg.com", "douyinpic.com", "douyinstatic.com")


class ProductReviewError(RuntimeError):
    pass


def utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def text(value, limit=20000):
    return str(value or "").strip()[:limit]


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def source_image(value):
    """Validate source evidence without persisting temporary query signatures."""
    try:
        parsed = urlsplit(text(value.get("url") if isinstance(value, dict) else value, 8192))
        host = (parsed.hostname or "").lower()
        if parsed.scheme != "https" or parsed.username or parsed.password:
            return None
        if not any(host == base or host.endswith("." + base) for base in ALLOWED_HOSTS):
            return None
        canonical = urlunsplit((parsed.scheme, parsed.netloc, parsed.path, "", ""))
        return {"identity": digest(canonical), "url": "" if parsed.query or parsed.fragment else canonical}
    except (ValueError, TypeError):
        return None


def normalize_part(raw):
    if isinstance(raw, str):
        raw = {"content": raw}
    if not isinstance(raw, dict):
        return None
    evidence = [source_image(value) for value in (raw.get("images") or [])[:60]]
    evidence = [value for value in evidence if value]
    evidence = list({value["identity"]: value for value in evidence}.values())
    content = text(raw.get("content"))
    if not content and not evidence:
        return None
    return {"content": content, "date_text": text(raw.get("dateText"), 180),
            "images": [value["url"] for value in evidence if value["url"]],
            "image_count": len(evidence), "omitted_image_link_count": sum(not value["url"] for value in evidence),
            "media_fingerprint": digest(sorted(value["identity"] for value in evidence))}


def normalize_review(raw, manifest):
    if not isinstance(raw, dict):
        return None
    part = normalize_part(raw)
    if not part:
        return None
    followups = [normalize_part(value) for value in (raw.get("followups") or [])[:30]]
    followups = [value for value in followups if value]
    reply = normalize_part(raw.get("merchantReply"))
    author = text(raw.get("reviewerName"), 300) or "匿名评价者"
    author_key = digest([manifest["task_id"], author])
    platform_id = text(raw.get("reviewId"), 160)
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,160}", platform_id):
        platform_id = ""
    sku = text(raw.get("purchasedSku"), 1200)
    # Late gallery/reply/helpful-count updates must upsert a text review,
    # not manufacture a second review. Media-only cards use their first
    # validated gallery asset as bounded fallback evidence (never its token).
    first_media = next((source_image(value) for value in (raw.get("images") or []) if source_image(value)), None)
    identity = [author_key, part["date_text"], sku, part["content"],
                "" if part["content"] else first_media["identity"] if first_media else ""]
    return {"review_id": platform_id or "derived_" + digest(identity),
            "id_source": "platform" if platform_id else "dom_fallback",
            "source_work_id": manifest["source_work_id"], "product_id": manifest["product"]["product_id"],
            "reviewer": author if manifest["privacy_mode"] == "raw" else "评论者_" + author_key[:12],
            **part, "purchased_sku": sku,
            "content_status": "text" if part["content"] else "media_only",
            "helpful_count": raw.get("helpfulCount") if isinstance(raw.get("helpfulCount"), int) and raw["helpfulCount"] >= 0 else None,
            "followups": followups, "merchant_reply": reply,
            "has_unparsed_followup": raw.get("hasUnparsedFollowup") is True,
            "collected_at": utc_now()}


def atomic_text(path, value):
    temporary = path.with_name(path.name + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write(value)
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temporary, path)


def save_rows(path, rows):
    atomic_text(path, "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows.values()))


def save_manifest(path, manifest):
    manifest["updated_at"] = utc_now()
    atomic_text(path, json.dumps(manifest, ensure_ascii=False, indent=2))


def build_product_review_workbook(out_dir, manifest, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    book = Workbook()
    info = book.active
    info.title = "采集说明"
    for row in [("商品", manifest["product"]["title"]), ("店铺", manifest["product"]["shop_name"]),
                ("来源作品ID", manifest["source_work_id"]), ("商品ID", manifest["product"]["product_id"]),
                ("状态", manifest["status"]), ("完整性", manifest["completeness"]),
                ("停止原因", manifest["done_reason"]), ("已保存去重评价", len(rows)),
                ("页面声明商品评价数", manifest.get("declared_review_count")),
                ("当前筛选", manifest.get("filter_label")),
                ("边界", "仅当前筛选下已读取评价；未下载评价图片或视频；签名链接不导出；不是平台内部全量。")]:
        info.append(row)
    sheet = book.create_sheet("商品评价")
    sheet.append(["评价内容", "评价者", "页面日期", "已购规格", "有用数", "内容类型", "页面图片数", "公开图片引用（未下载）",
                  "未导出签名链接数", "追评", "商家回复", "评价ID", "ID来源", "来源作品ID", "商品ID", "采集时间"])
    for value in rows.values():
        sheet.append([value["content"], value["reviewer"], value["date_text"], value["purchased_sku"], value["helpful_count"],
                      value["content_status"], value["image_count"], "\n".join(value["images"]), value["omitted_image_link_count"],
                      json.dumps(value["followups"], ensure_ascii=False), json.dumps(value["merchant_reply"], ensure_ascii=False) if value["merchant_reply"] else "",
                      value["review_id"], value["id_source"], value["source_work_id"], value["product_id"], value["collected_at"]])
    for page in book:
        page.freeze_panes = "A2"
        if page == sheet:
            page.auto_filter.ref = page.dimensions
        for row in page:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, str):
                    # Source prose is data, never an Excel formula.
                    cell.data_type = "s"
        for cell in page[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
        for column in page.columns:
            page.column_dimensions[column[0].column_letter].width = 30 if page == info else 24
    sheet.column_dimensions["A"].width = 52
    sheet.column_dimensions["H"].width = 56
    info.column_dimensions["B"].width = 80
    target = Path(out_dir) / "商品评价.xlsx"
    temporary = target.with_name("商品评价.tmp.xlsx")
    book.save(temporary)
    os.replace(temporary, target)
    return target


def _call(page, method, argument=None):
    return page.evaluate(f"value => {BRIDGE}.{method}(value)", argument)


def valid_lease(value):
    return isinstance(value, dict) and all(isinstance(value.get(key), str) and value[key] for key in LEASE_FIELDS)


def safe_product(value):
    value = value if isinstance(value, dict) else {}
    product_id = text(value.get("productId") or value.get("product_id"), 30)
    return {"product_id": product_id if re.fullmatch(r"\d{6,24}", product_id) else "",
            "title": text(value.get("title"), 300), "shop_name": text(value.get("shopName") or value.get("shop_name"), 100)}


def validate_resume(previous, surface, expected_work_id, privacy_mode):
    if previous.get("source_work_id") != expected_work_id or previous.get("privacy_mode") != privacy_mode:
        raise ProductReviewError("续跑作品或隐私模式不同；旧结果未修改，请使用新目录")
    product = safe_product(surface.get("product"))
    if previous.get("product") != product or not valid_lease(surface.get("lease")):
        raise ProductReviewError("商品身份不一致；旧结果未修改，请使用新目录")
    old_lease = previous.get("lease") or {}
    if old_lease.get("filterKey") != surface["lease"]["filterKey"]:
        raise ProductReviewError("商品评价筛选已变化；旧结果未修改，请使用新目录")
    if not product["product_id"] and old_lease != surface["lease"]:
        raise ProductReviewError("未取得可靠商品ID，不能跨页面或会话合并评价；旧结果未修改，请使用新目录")


def _finish_manifest(manifest, message, rows):
    complete = message.get("status") == "finished" and message.get("completeness") == COMPLETE and message.get("exhausted") is True
    empty = message.get("emptyConfirmed") is True and message.get("visibleReviewCount") == 0
    complete = complete and (bool(rows) or empty) and not any(row.get("has_unparsed_followup") for row in rows.values())
    manifest.update(status="complete" if complete else "paused" if message.get("status") == "paused" else "partial",
                    completeness=COMPLETE if complete else text(message.get("completeness"), 100) if text(message.get("completeness")).startswith("partial_") else "partial_unverified_end",
                    done_reason=text(message.get("doneReason"), 100), finished_at=utc_now(), exhausted=complete)


def collect_product_reviews(page, expected_work_id, expected_product, out_dir, *, max_reviews=200, max_scrolls=200,
                            max_seconds=600, privacy_mode="hash", resume=False) -> int:
    """Collect one opened product's current-filter reviews; 0 complete, 3 partial.

    Write product_review_control.json with {"action":"pause"} to stop after
    the in-flight batch is durably saved. Never start a second worker to poll.
    """
    expected_work_id = str(expected_work_id)
    product = safe_product(expected_product)
    if not re.fullmatch(r"\d+", expected_work_id) or not product["title"] or privacy_mode not in {"hash", "raw"}:
        raise ProductReviewError("商品评价要求明确作品、商品标题与合法隐私模式")
    if not 1 <= max_reviews <= 2000 or not 1 <= max_scrolls <= 200 or not 1 <= max_seconds <= 600:
        raise ProductReviewError("商品评价限额须为1–2000条、1–200次滚动、1–600秒")
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest_path, rows_path = out_dir / "product_review_manifest.json", out_dir / "product_reviews.jsonl"
    lock_path = out_dir / ".product_reviews.lock"
    try:
        lock = lock_path.open("x", encoding="utf-8")
    except FileExistsError as error:
        raise ProductReviewError("该目录存在商品评价任务锁；先确认原进程结束，不要并行重跑") from error
    lock.write(str(os.getpid())); lock.close()
    manifest = None
    rows = {}
    try:
        if (manifest_path.exists() or rows_path.exists()) and not resume:
            raise ProductReviewError("输出目录已有商品评价，使用resume或新目录；旧文件未修改")
        previous = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else None
        if resume and (not previous or not rows_path.exists()):
            raise ProductReviewError("缺少可核对的商品评价断点；请使用新目录")
        if not previous:
            manifest = {"schema_version": "1.0", "kind": "douyin_product_reviews", "task_id": uuid.uuid4().hex,
                        "run_id": uuid.uuid4().hex, "source_work_id": expected_work_id, "product": product,
                        "lease": None, "privacy_mode": privacy_mode, "status": "partial", "completeness": "partial_preparing",
                        "done_reason": "preparing", "review_count": 0, "progress": {}, "started_at": utc_now()}
        for name in ("douyin_product_review_collector.js", "product_review_page_bridge.js"):
            page.evaluate((ROOT / name).read_text(encoding="utf-8"))
        surface = _call(page, "configure", {"expectedWorkId": expected_work_id,
                        "product": {"title": product["title"], "shopName": product["shop_name"], "productId": product["product_id"]},
                        "limits": {"maxRows": max_reviews, "maxScrolls": max_scrolls, "maxMs": max_seconds * 1000}})
        if not surface.get("lease"):
            surface = _call(page, "open")
        if previous:
            validate_resume(previous, surface, expected_work_id, privacy_mode)
            for line in rows_path.read_text(encoding="utf-8").splitlines():
                row = json.loads(line)
                rows[row["id_source"] + ":" + row["review_id"]] = row
            if previous.get("status") == "complete" and previous.get("completeness") == COMPLETE and previous.get("exhausted") is True:
                if previous.get("review_count") != len(rows):
                    raise ProductReviewError("商品评价断点数量与文件不符；旧结果未修改，请核对后重试")
                if not (out_dir / "商品评价.xlsx").exists():
                    build_product_review_workbook(out_dir, previous, rows)
                return 0
        manifest = {"schema_version": "1.0", "kind": "douyin_product_reviews", "task_id": previous["task_id"] if previous else manifest["task_id"],
                    "run_id": uuid.uuid4().hex, "source_work_id": expected_work_id, "product": safe_product(surface.get("product")),
                    "lease": surface.get("lease"), "filter_label": text(surface.get("filterLabel"), 1000), "privacy_mode": privacy_mode,
                    "status": "running", "completeness": "partial_reading", "done_reason": "", "started_at": utc_now(),
                    "review_count": len(rows), "visible_review_count": surface.get("visibleReviewCount"),
                    "declared_review_count": surface.get("declaredReviewCount"), "progress": {},
                    "files": {"rows": rows_path.name, "workbook": "商品评价.xlsx"},
                    "boundaries": ["current_filter_only", "no_review_media_download", "temporary_image_links_omitted"]}
        save_rows(rows_path, rows)
        save_manifest(manifest_path, manifest)
        if not surface.get("ready") or not valid_lease(surface.get("lease")):
            manifest.update(status="partial", completeness="partial_surface_unverified", done_reason=text(surface.get("reasonCode"), 100))
            save_manifest(manifest_path, manifest)
            build_product_review_workbook(out_dir, manifest, rows)
            return 3
        if len(rows) >= max_reviews:
            manifest.update(status="partial", completeness="partial_limit_sample", done_reason="run_budget")
            save_manifest(manifest_path, manifest)
            build_product_review_workbook(out_dir, manifest, rows)
            return 3
        _call(page, "start", {"id": manifest["task_id"], "runId": manifest["run_id"], "lease": manifest["lease"], "reviewCount": len(rows)})
        paused = False
        stop_started = None
        began = time.monotonic()
        ended = False
        limit_reached = False
        while not ended:
            try:
                control = out_dir / "product_review_control.json"
                should_pause = control.exists() and json.loads(control.read_text(encoding="utf-8")).get("action") == "pause"
                if not paused and (should_pause or time.monotonic() - began > max_seconds + 5):
                    _call(page, "pause"); paused = True; stop_started = time.monotonic()
                polled = _call(page, "poll")
                for item in polled.get("items") or []:
                    message = item["message"]
                    if message.get("taskId") != manifest["task_id"] or message.get("runId") != manifest["run_id"] or message.get("lease") != manifest["lease"]:
                        _call(page, "ack", {"id": item["id"], "response": {"ok": False, "ignored": True}})
                        continue
                    kind = message.get("type")
                    if kind == "CAPTURE_DOUYIN_COMMERCE_REVIEWS":
                        normalized = [normalize_review(row, manifest) for row in message.get("rows") or []]
                        if any(row is None for row in normalized):
                            _call(page, "ack", {"id": item["id"], "response": {"ok": False, "error": "invalid_review_payload"}})
                            continue
                        for row in normalized:
                            key = row["id_source"] + ":" + row["review_id"]
                            if key not in rows and len(rows) >= max_reviews:
                                _call(page, "pause"); paused = True; stop_started = time.monotonic()
                                limit_reached = True
                                break
                            row["collected_at"] = rows.get(key, row)["collected_at"]
                            rows[key] = row
                        save_rows(rows_path, rows)
                    progress = message.get("progress") or {}
                    if progress.get("sequence", -1) > manifest["progress"].get("sequence", -1):
                        manifest["progress"] = {key: value for key, value in progress.items() if key in
                            {"sequence", "phase", "elapsedMs", "scrollActions", "visibleReviewCount", "parsedReviewCount", "unparsedReviewCount"}}
                    for source, target in (("visibleReviewCount", "visible_review_count"), ("declaredReviewCount", "declared_review_count")):
                        if isinstance(message.get(source), int) and message[source] >= 0:
                            manifest[target] = message[source]
                    manifest["review_count"] = len(rows)
                    if kind == "FINISH_DOUYIN_COMMERCE_REVIEWS":
                        _finish_manifest(manifest, message, rows); ended = True
                        if limit_reached:
                            manifest.update(status="partial", completeness="partial_limit_sample", done_reason="run_budget", exhausted=False)
                    save_manifest(manifest_path, manifest)
                    _call(page, "ack", {"id": item["id"], "response": {"ok": True, "task": {"id": manifest["task_id"], "runId": manifest["run_id"], "reviewCount": len(rows)}}})
                if not ended and not polled.get("active"):
                    manifest.update(status="partial", completeness="partial_collector_interrupted", done_reason="collector_interrupted"); ended = True
                if stop_started and time.monotonic() - stop_started > 15:
                    raise ProductReviewError("暂停未完成确认；已落盘结果保留，禁止继续导航或新建采集")
                if not ended:
                    page.wait_for_timeout(150)
            except KeyboardInterrupt:
                if paused:
                    raise ProductReviewError("二次中断发生于保存确认期间；结果保留，停止使用当前采集页面")
                _call(page, "pause"); paused = True; stop_started = time.monotonic()
        save_manifest(manifest_path, manifest)
        build_product_review_workbook(out_dir, manifest, rows)
        return 0 if manifest["status"] == "complete" else 3
    except ProductReviewError:
        if manifest:
            manifest.update(status="partial", completeness="partial_interrupted", done_reason="interrupted")
            save_rows(rows_path, rows); save_manifest(manifest_path, manifest)
            build_product_review_workbook(out_dir, manifest, rows)
        raise
    except Exception as error:
        if manifest:
            try:
                _call(page, "pause")
            except Exception:
                pass
            manifest.update(status="partial", completeness="partial_browser_error", done_reason="browser_or_storage_error")
            save_rows(rows_path, rows); save_manifest(manifest_path, manifest)
            build_product_review_workbook(out_dir, manifest, rows)
        raise ProductReviewError("商品评价读取中断；已保存结果保留，请检查原页面后再续跑") from error
    finally:
        lock_path.unlink(missing_ok=True)
